import asyncio
import hashlib
import io
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urljoin, urlparse

import httpx
from bs4 import BeautifulSoup
from docx import Document
from pypdf import PdfReader

from .config import get_settings


HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) SBER-Hackathon-RAG/1.0"
}


def clean_text(text: str) -> str:
    text = text.replace("\xa0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


_CHROME_TAGS = ("script", "style", "noscript", "svg", "nav", "footer", "header", "aside")
_CHROME_ATTR_RE = re.compile(
    r"\b(menu|navbar|nav-menu|main-menu|sidebar|breadcrumb|site-header|site-footer|top-menu|header|footer|nav)\b",
    re.IGNORECASE,
)
_MENU_LINES = frozenset({
    "главная",
    "расчет жку",
    "образец формы единого платежного документа",
    "пункты приема платежей за жку",
    "наименование получателей платежей без комиссии",
    "наименование получателей платежей с комиссией",
    "вопрос-ответ",
    "транспортная карта",
    "порядок вывода банковской карты из стоп-листа",
    "транспортная карта «тройка»",
    "транспортная карта тройка",
    "социальная транспортная карта (стк)",
    "правила транспортной системы «сбертройка»",
    "о компании",
    "центры обслуживания населения",
    "нормативно-правовые акты",
    "документы общества",
    "раскрытие информации",
    "принципалы",
    "руководство",
    "вакансии",
    "противодействие коррупции",
    "новости",
    "контакты",
    "личный кабинет",
    "личный кабинет жкх",
    "личный кабинет пассажира",
    "тула",
    "орел",
    "орёл",
    "серпухов",
    "воронеж",
    "сайт \"сбербилет\"",
    "сайт \"нспк\"",
    "сайт «сбербилет»",
    "сайт «нспк»",
})
_HEADER_LINES = frozenset({
    "оеирц",
    "«областной единый информационно-расчетный центр»",
    "областной единый информационно-расчетный центр",
    "акционерное общество",
})
_HEADING_LINE_RE = re.compile(r"^\d+(?:\.\d+){1,4}\.?\s+\S")
_URL_OR_WORD_RE = re.compile(r"https?://\S+|\S+")


def _norm_line(line: str) -> str:
    normalized = line.replace("ё", "е").replace("Ё", "Е")
    normalized = normalized.strip().strip("«»\"'").lower()
    return re.sub(r"\s+", " ", normalized)


def _is_menu_line(line: str) -> bool:
    return _norm_line(line) in {_norm_line(item) for item in _MENU_LINES}


def strip_site_chrome(text: str) -> str:
    """Drop the OEIRC masthead and a run of navigation lines, keep body sentences."""
    lines = text.split("\n")
    start = 0
    skipped = 0
    while start < len(lines) and skipped < 6:
        key = _norm_line(lines[start])
        if not key:
            start += 1
            continue
        if key not in {_norm_line(item) for item in _HEADER_LINES}:
            break
        start += 1
        skipped += 1
    lines = lines[start:]

    kept: list[str] = []
    index = 0
    while index < len(lines):
        if _is_menu_line(lines[index]):
            end = index
            while end < len(lines) and (not _norm_line(lines[end]) or _is_menu_line(lines[end])):
                end += 1
            menu_count = sum(1 for line in lines[index:end] if _is_menu_line(line))
            if menu_count >= 4:
                index = end
                continue
        kept.append(lines[index])
        index += 1
    return clean_text("\n".join(kept))


def _is_navigation_block(tag) -> bool:
    if tag.name in {"body", "html", "[document]"}:
        return False
    # OEIRC wraps the menu and the article in one table. Drop only the menu itself.
    if tag.get("id") == "content" or tag.find(id="content") is not None:
        return False
    identity = " ".join(
        [
            tag.get("id") or "",
            " ".join(tag.get("class") or []),
            tag.get("role") or "",
        ]
    )
    if _CHROME_ATTR_RE.search(identity):
        return True
    links = tag.find_all("a")
    if len(links) < 4:
        return False
    link_texts = [" ".join(link.stripped_strings) for link in links]
    menu_hits = sum(1 for item in link_texts if _is_menu_line(item))
    return menu_hits >= 4


def _strip_chrome(soup: BeautifulSoup) -> None:
    for tag in soup.find_all(_CHROME_TAGS):
        tag.decompose()
    candidates = [tag for tag in soup.find_all(True) if _is_navigation_block(tag)]
    candidate_ids = {id(tag) for tag in candidates}
    for tag in candidates:
        if any(id(parent) in candidate_ids for parent in tag.parents):
            continue
        tag.decompose()


def _is_heading_line(line: str) -> bool:
    if line.startswith("# "):
        return True
    if len(line) > 100:
        return False
    return bool(_HEADING_LINE_RE.match(line))


def extract_html(content: bytes) -> str:
    soup = BeautifulSoup(content, "html.parser")
    _strip_chrome(soup)
    blocks: list[str] = []
    root = soup.find(id="content") or soup.body or soup
    for element in root.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "table"]):
        if element.name != "table" and element.find_parent("table"):
            continue
        if element.name in {"p", "li"} and element.find_parent(["li", "p"]):
            continue
        if element.name == "table":
            if element.find(["p", "h1", "h2", "h3", "h4", "h5", "h6", "table", "nav"]) is not None:
                continue
            rows: list[str] = []
            for row in element.find_all("tr"):
                cells = [cell.get_text(" ", strip=True) for cell in row.find_all(["th", "td"])]
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                blocks.append("\n".join(rows))
            continue
        text = " ".join(element.stripped_strings).strip()
        if not text:
            continue
        if element.name.startswith("h"):
            blocks.append(f"# {text}")
        else:
            blocks.append(text)
    plain = clean_text(root.get_text("\n"))
    structured = "\n\n".join(blocks)
    if len(structured) < len(plain) * 0.5:
        return strip_site_chrome(plain)
    return strip_site_chrome(structured)


def extract_docx(content: bytes) -> str:
    doc = Document(io.BytesIO(content))
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    return clean_text("\n".join(lines))


def extract_pdf(content: bytes) -> str:
    reader = PdfReader(io.BytesIO(content))
    return clean_text("\n".join((page.extract_text() or "") for page in reader.pages))


def extract_text(content: bytes, kind: str) -> str:
    kind = kind.lower()
    if kind == "docx":
        return extract_docx(content)
    if kind == "pdf":
        return extract_pdf(content)
    return extract_html(content)


def split_paragraph(paragraph: str, chunk_size: int) -> list[str]:
    """Break a long paragraph on whitespace. Words and URLs stay intact."""
    paragraph = paragraph.strip()
    if not paragraph or len(paragraph) <= chunk_size:
        return [paragraph] if paragraph else []
    tokens = _URL_OR_WORD_RE.findall(paragraph)
    parts: list[str] = []
    current: list[str] = []
    current_len = 0
    for token in tokens:
        extra = len(token) if not current else len(token) + 1
        if current and current_len + extra > chunk_size:
            parts.append(" ".join(current))
            current = [token]
            current_len = len(token)
            continue
        current.append(token)
        current_len += extra
    if current:
        parts.append(" ".join(current))
    return parts


def _iter_blocks(text: str) -> list[tuple[str, str | list[str]]]:
    blocks: list[tuple[str, str | list[str]]] = []
    paragraph: list[str] = []
    table: list[str] = []

    def flush_paragraph() -> None:
        nonlocal paragraph
        body = "\n".join(paragraph).strip()
        paragraph = []
        if body:
            blocks.append(("paragraph", body))

    def flush_table() -> None:
        nonlocal table
        if table:
            blocks.append(("table", list(table)))
            table = []

    for raw in text.split("\n"):
        line = raw.strip()
        if not line:
            flush_table()
            flush_paragraph()
            continue
        if " | " in line and not line.startswith("#"):
            flush_paragraph()
            table.append(line)
            continue
        flush_table()
        if _is_heading_line(line):
            flush_paragraph()
            heading = line if line.startswith("#") else f"# {line}"
            blocks.append(("heading", heading))
            continue
        paragraph.append(line)
    flush_table()
    flush_paragraph()
    return blocks


def _chunk_table(rows: list[str], chunk_size: int) -> list[str]:
    if not rows:
        return []
    header, data = rows[0], rows[1:]
    if not data:
        return [header]
    chunks: list[str] = []
    group: list[str] = []
    for row in data:
        candidate = "\n".join([header, *group, row])
        if group and len(candidate) > chunk_size:
            chunks.append("\n".join([header, *group]))
            group = [row]
            continue
        group.append(row)
    if group:
        chunks.append("\n".join([header, *group]))
    return chunks


def _join_parts(parts: list[str]) -> str:
    return "\n\n".join(parts).strip()


def is_template_fragment(text: str) -> bool:
    """SPA placeholders are not a schedule and must not enter the index."""
    return "{{" in text or "Загрузка маршрутов" in text


def chunk_text(text: str, chunk_size: int = 1500, overlap: int = 220) -> list[str]:
    """Split by headings and paragraphs. Overlap is one whole paragraph, never a character tail."""
    use_overlap = overlap > 0
    chunks: list[str] = []
    current: list[str] = []
    pending_heading = ""
    last_complete_paragraph = ""

    def flush(*, keep_overlap: bool) -> str:
        nonlocal current, last_complete_paragraph
        if not current:
            return ""
        chunks.append(_join_parts(current))
        overlap_paragraph = last_complete_paragraph if keep_overlap and use_overlap else ""
        current = []
        if not keep_overlap:
            last_complete_paragraph = ""
        return overlap_paragraph

    def start_with(overlap_paragraph: str, parts: list[str]) -> None:
        nonlocal current
        current = []
        if overlap_paragraph and len(overlap_paragraph) + 2 + len(_join_parts(parts)) <= chunk_size:
            if overlap_paragraph not in parts:
                current.append(overlap_paragraph)
        current.extend(parts)

    for kind, payload in _iter_blocks(strip_site_chrome(text)):
        if kind == "heading":
            flush(keep_overlap=False)
            pending_heading = str(payload)
            continue
        if kind == "table":
            flush(keep_overlap=False)
            pending_heading = ""
            chunks.extend(_chunk_table(list(payload), chunk_size))
            continue

        pieces = split_paragraph(str(payload), chunk_size)
        complete = pieces[0] if len(pieces) == 1 else ""
        for index, piece in enumerate(pieces):
            prefix = [pending_heading] if pending_heading else []
            pending_heading = ""
            addition = prefix + [piece]
            candidate = _join_parts(current + addition)
            if current and len(candidate) > chunk_size:
                carried = flush(keep_overlap=True)
                start_with(carried, addition)
            else:
                current.extend(addition)
            if complete and index == 0:
                last_complete_paragraph = complete
            else:
                last_complete_paragraph = ""
    if pending_heading:
        current.append(pending_heading)
    flush(keep_overlap=False)
    return [chunk for chunk in chunks if chunk and not _is_heading_only(chunk)]


def _is_heading_only(chunk: str) -> bool:
    lines = [line.strip() for line in chunk.splitlines() if line.strip()]
    return bool(lines) and all(line.startswith("# ") for line in lines)


def empty_source_message(title: str, kind: str) -> str:
    if (kind or "").lower() == "doc":
        return f"[INFO] {title}: формат .doc без извлекаемого текста, в корпус не пишется"
    return f"[INFO] {title}: downloaded only ({kind})"


def safe_filename(source_id: str, kind: str, url: str) -> str:
    suffix = {"html": ".html", "pdf": ".pdf", "docx": ".docx", "doc": ".doc"}.get(kind.lower())
    if not suffix:
        suffix = Path(urlparse(url).path).suffix or ".bin"
    return f"{re.sub(r'[^a-zA-Z0-9._-]+', '_', source_id)}{suffix}"


PASSENGER_DISCOVERY_PRIORITY = 68
OPERATOR_TECHNICAL_TITLE_MARKERS = ("технический регламент", "участников асоп")


def classify_document_audience(title: str) -> str:
    """Passenger rules stay searchable; operator manuals stay out of that index."""
    haystack = title.lower()
    if any(marker in haystack for marker in OPERATOR_TECHNICAL_TITLE_MARKERS):
        return "operator_technical"
    return "passenger"


def file_content_hash(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def annotate_discovered_document(doc: dict) -> dict:
    audience = classify_document_audience(doc.get("title", ""))
    doc["scope"] = "region"
    doc["audience"] = audience
    doc["doc_type"] = doc.get("kind") or "pdf"
    doc["priority"] = min(int(doc.get("priority", PASSENGER_DISCOVERY_PRIORITY)), PASSENGER_DISCOVERY_PRIORITY)
    doc["current"] = True
    doc["enabled"] = audience == "passenger"
    doc["discovered"] = True
    return doc


def catalog_entry_for_discovered(source: dict, content: bytes, fetched_at: str | None = None) -> dict:
    annotated = annotate_discovered_document(dict(source))
    annotated["content_hash"] = file_content_hash(content)
    entry = {
        "id": annotated["id"],
        "scope": annotated["scope"],
        "audience": annotated["audience"],
        "doc_type": annotated["doc_type"],
        "title": annotated.get("title"),
        "url": annotated.get("url"),
        "kind": annotated.get("kind"),
        "authority": annotated.get("authority"),
        "category": annotated.get("category"),
        "priority": annotated["priority"],
        "current": annotated["current"],
        "enabled": annotated["enabled"],
        "content_hash": annotated["content_hash"],
        "discovered": True,
    }
    if fetched_at:
        entry["fetched_at"] = fetched_at
    return entry


def source_freshness(loaded: dict) -> dict:
    """Hash and capture time stored on both the catalog row and each chunk."""
    freshness = {}
    if loaded.get("content_hash"):
        freshness["content_hash"] = loaded["content_hash"]
    if loaded.get("fetched_at"):
        freshness["fetched_at"] = loaded["fetched_at"]
    return freshness


def apply_source_freshness(catalog: list[dict], freshness_by_id: dict[str, dict]) -> list[dict]:
    updated = []
    for row in catalog:
        stamp = freshness_by_id.get(row.get("id"))
        if not stamp:
            updated.append(row)
            continue
        merged = dict(row)
        merged.update({key: value for key, value in stamp.items() if value})
        updated.append(merged)
    return updated


def chunk_record(source: dict, loaded: dict, part_index: int, text: str) -> dict:
    return {
        "id": f"{source['id']}:{part_index}",
        "source_id": source["id"],
        "title": source.get("title"),
        "url": source.get("url"),
        "resolved_url": loaded.get("resolved_url"),
        "authority": source.get("authority"),
        "category": source.get("category"),
        "priority": source.get("priority", 50),
        "current": source.get("current", True),
        "fetched_at": loaded.get("fetched_at"),
        "content_hash": loaded.get("content_hash"),
        "text": text,
    }


def fare_amount_conflicts(sources: list[dict], cards: list[dict]) -> list[dict]:
    """Current sources of one category whose fare cards disagree on amounts.

    Priority stays a human field. This only reports the disagreement.
    """
    by_id = {row["id"]: row for row in sources if row.get("id")}
    grouped: dict[str, dict[str, set]] = {}
    for card in cards:
        source = by_id.get(card.get("source_id"))
        if not source or source.get("current") is not True:
            continue
        category = source.get("category")
        amount = card.get("amount")
        if not category or amount is None:
            continue
        grouped.setdefault(str(category), {}).setdefault(source["id"], set()).add(amount)
    conflicts = []
    for category, by_source in grouped.items():
        if len(by_source) < 2:
            continue
        amounts: set = set()
        for values in by_source.values():
            amounts.update(values)
        if len(amounts) < 2:
            continue
        conflicts.append({
            "category": category,
            "source_ids": sorted(by_source),
            "amounts": sorted(amounts, key=lambda value: (str(type(value)), value)),
        })
    return conflicts


def warn_conflicting_fare_amounts(sources: list[dict], cards: list[dict]) -> list[dict]:
    """Print a warning and keep every card. Do not pick a winner or edit priority."""
    for conflict in fare_amount_conflicts(sources, cards):
        source_ids = ", ".join(conflict["source_ids"])
        amounts = ", ".join(str(amount) for amount in conflict["amounts"])
        print(
            f"[WARN] category {conflict['category']}: current sources {source_ids} "
            f"contain different fare amounts ({amounts}); both kept, priority unchanged"
        )
    return list(cards)


def dedupe_discovered_sources(sources: list[dict]) -> list[dict]:
    seen_hashes: set[str] = set()
    unique: list[dict] = []
    for source in sources:
        digest = source.get("content_hash")
        if digest:
            if digest in seen_hashes:
                continue
            seen_hashes.add(digest)
        unique.append(source)
    return unique


def merge_discovered_catalog(existing: list[dict], discovered: list[dict]) -> list[dict]:
    curated = [row for row in existing if not row.get("discovered")]
    return curated + dedupe_discovered_sources(discovered)


def discover_documents(base_url: str, content: bytes) -> list[dict]:
    soup = BeautifulSoup(content, "html.parser")
    docs: list[dict] = []
    seen: set[str] = set()
    for link in soup.find_all("a", href=True):
        href = str(link["href"]).strip()
        absolute = urljoin(base_url, href)
        path = urlparse(absolute).path.lower()
        if not path.endswith((".pdf", ".docx", ".doc")):
            continue
        label = " ".join(link.stripped_strings).strip() or Path(path).name
        haystack = f"{label} {absolute}".lower()
        transport_terms = (
            "транспорт", "тройк", "проезд", "льгот", "маршрут", "карт", "асоп",
            "билет", "пассажир", "перевоз", "измен", "асоп", "ткп",
            "83", "59", "412", "3661"
        )
        if not any(term in haystack for term in transport_terms):
            continue
        if absolute in seen:
            continue
        seen.add(absolute)
        if path.endswith(".pdf"):
            kind = "pdf"
        elif path.endswith(".docx"):
            kind = "docx"
        else:
            kind = "doc"
        digest = hashlib.sha1(absolute.encode("utf-8")).hexdigest()[:10]
        docs.append({
            "id": f"oeirc-discovered-{digest}",
            "title": f"ОЕИРЦ — {label}",
            "url": absolute,
            "kind": kind,
            "authority": "АО ОЕИРЦ",
            "category": "нормативные_документы",
            "priority": 68,
            "current": True,
            "enabled": True,
            "discovered": True,
        })
    return docs


def discover_transport_documents(content: bytes, base_url: str, parent: dict) -> list[dict]:
    """Backwards-compatible, testable entry point with parent metadata."""
    docs = discover_documents(base_url, content)
    for doc in docs:
        doc["authority"] = parent.get("authority", doc["authority"])
        doc["priority"] = min(int(parent.get("priority", 75)), PASSENGER_DISCOVERY_PRIORITY)
        annotate_discovered_document(doc)
    return docs


async def fetch_source(client: httpx.AsyncClient, source: dict, raw_dir: Path) -> tuple[dict, list[dict]]:
    response = await client.get(source["url"], headers=HEADERS, follow_redirects=True)
    response.raise_for_status()

    raw_dir.mkdir(parents=True, exist_ok=True)
    raw_path = raw_dir / safe_filename(source["id"], source.get("kind", "html"), str(response.url))
    raw_path.write_bytes(response.content)

    kind = source.get("kind", "html")
    if kind == "doc":
        text = ""
    else:
        text = extract_text(response.content, kind)

    discovered = []
    if source.get("discover_documents") and kind == "html":
        discovered = discover_transport_documents(response.content, str(response.url), source)

    loaded = {
        **source,
        "resolved_url": str(response.url),
        "fetched_at": datetime.now(timezone.utc).isoformat(),
        "http_status": response.status_code,
        "raw_file": str(raw_path),
        "content_hash": file_content_hash(response.content),
        "text": text,
    }
    return loaded, discovered


async def build_chunks() -> list[dict]:
    settings = get_settings()
    sources = json.loads(settings.sources_path.read_text(encoding="utf-8"))
    enabled = [s for s in sources if s.get("enabled", True) and s.get("current", True)]
    results: list[dict] = []
    snapshots: list[dict] = []
    discovered_catalog: list[dict] = []
    freshness_by_id: dict[str, dict] = {}
    seen_content_hashes: set[str] = set()
    raw_dir = settings.chunks_path.parent / "raw_sources"

    async with httpx.AsyncClient(timeout=45, verify=False, follow_redirects=True) as client:
        queue = list(enabled)
        seen_ids = {s["id"] for s in queue}
        index = 0
        while index < len(queue):
            source = queue[index]
            index += 1
            try:
                loaded, discovered = None, []
                last_error: Exception | None = None
                for attempt in range(3):
                    try:
                        loaded, discovered = await fetch_source(client, source, raw_dir)
                        last_error = None
                        break
                    except Exception as exc:
                        last_error = exc
                        if attempt == 2:
                            raise
                if last_error is not None or loaded is None:
                    raise last_error or RuntimeError("source was not fetched")
                for doc in discovered:
                    if doc["id"] not in seen_ids:
                        seen_ids.add(doc["id"])
                        queue.append(doc)

                freshness_by_id[source["id"]] = source_freshness(loaded)
                if source.get("discovered"):
                    digest = loaded["content_hash"]
                    if digest in seen_content_hashes:
                        print(f"[INFO] {source['title']}: duplicate content, catalog keeps one source")
                        continue
                    seen_content_hashes.add(digest)
                    entry = catalog_entry_for_discovered(
                        source, Path(loaded["raw_file"]).read_bytes(), loaded.get("fetched_at")
                    )
                    discovered_catalog.append(entry)
                    source = {**source, **entry}
                    if not entry["enabled"]:
                        snapshots.append({**loaded, **entry, "text": ""})
                        print(f"[INFO] {source['title']}: {entry['audience']}, skipped in passenger index")
                        continue

                snapshots.append(loaded)

                if not loaded["text"]:
                    print(empty_source_message(source["title"], source.get("kind", "")))
                    continue

                parts = [part for part in chunk_text(loaded["text"]) if not is_template_fragment(part)]
                for part_index, text in enumerate(parts):
                    results.append(chunk_record(source, loaded, part_index, text))
                print(f"[OK] {source['title']}: {len(parts)} chunks")
            except Exception as exc:
                print(f"[WARN] {source['title']}: {exc!r}")

    if not results:
        raise RuntimeError("No RAG chunks were built. Check source availability/network.")

    catalog = apply_source_freshness(merge_discovered_catalog(sources, discovered_catalog), freshness_by_id)
    fare_cards_path = settings.chunks_path.parent / "fare_cards.json"
    if fare_cards_path.exists():
        fare_cards = json.loads(fare_cards_path.read_text(encoding="utf-8"))
        if isinstance(fare_cards, list):
            warn_conflicting_fare_amounts(catalog, fare_cards)
    settings.sources_path.write_text(
        json.dumps(catalog, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    settings.chunks_path.write_text(
        json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    snapshot_path = settings.chunks_path.parent / "source_snapshots.json"
    snapshot_path.write_text(
        json.dumps(snapshots, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(f"Saved {len(results)} chunks to {settings.chunks_path}")
    print(f"Saved {len(snapshots)} source snapshots to {snapshot_path}")
    print(f"Raw source copies: {raw_dir}")
    return results


async def main() -> None:
    chunks = await build_chunks()
    print(f"Agentic text-search corpus is ready: {len(chunks)} chunks")


if __name__ == "__main__":
    asyncio.run(main())
